"""Build a formatted Excel workbook from a test-case CSV (one sheet per module + Summary).

Usage:
    python scripts/build_testcases_excel.py <cases.csv> <output.xlsx> \
        --phase "R&D Phase 2" --notes [notes.json] --sources "source1; source2"

The CSV must use the standard 10-column layout:
    TC-ID, Title, Module, Precondition, Test Data, Steps, Expected Result, Actual Result, Status, Priority

The optional notes JSON maps TC-ID -> Notes (extra column appended in the sheets).
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER = ["TC-ID", "Title", "Module", "Precondition", "Test Data", "Steps",
          "Expected Result", "Actual Result", "Status", "Priority"]
WIDTHS = [22, 46, 18, 44, 44, 72, 72, 14, 10, 10]
NOTES_WIDTH = 52
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def read_cases(csv_path: Path) -> list[list[str]]:
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if rows[0] != HEADER:
        sys.exit(f"Unexpected header in {csv_path}: {rows[0]}")
    body = [r for r in rows[1:] if any(cell.strip() for cell in r)]
    for r in body:
        if len(r) != len(HEADER):
            sys.exit(f"Row does not have {len(HEADER)} columns: {r[0]}")
    return body


def write_sheet(ws, cases: list[list[str]], notes: dict[str, str]) -> None:
    ws.append(HEADER + (["Notes"] if notes else []))
    for row in cases:
        ws.append(row + ([notes.get(row[0], "")] if notes else []))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    widths = WIDTHS + ([NOTES_WIDTH] if notes else [])
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_summary(ws, body: list[list[str]], phase: str, sources: list[str],
                  summary_note: str) -> None:
    modules: dict[str, list[list[str]]] = {}
    for r in body:
        modules.setdefault(r[2], []).append(r)
    ws.append([f"Finalized Test Cases - {phase}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Total cases", len(body)])
    ws.append(["Generation date", date.today().isoformat()])
    for s in sources:
        ws.append(["Source", s])
    ws.append([])
    ws.append(["Module", "Cases", "High", "Medium", "Low"])
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    for mod in sorted(modules):
        cases = modules[mod]
        counts = {p: sum(1 for c in cases if c[9] == p) for p in ("High", "Medium", "Low")}
        ws.append([mod, len(cases), counts["High"], counts["Medium"], counts["Low"]])
        for cell in ws[ws.max_row]:
            cell.border = BORDER
    ws.append([])
    ws.append([summary_note])
    for i, width in enumerate([28, 12, 10, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("csv")
    parser.add_argument("output")
    parser.add_argument("--phase", required=True)
    parser.add_argument("--notes")
    parser.add_argument("--sources", default="")
    parser.add_argument(
        "--summary-note",
        default="Notes: PPT deck (Phase 2 Master V.2.0) is authoritative; SRS docx use cases used to fill detail gaps. "
                "Cells in the Notes column mark SRS/PPT conflicts and SRS-only additions.",
    )
    args = parser.parse_args()

    body = read_cases(Path(args.csv))
    notes = {}
    if args.notes:
        notes = json.loads(Path(args.notes).read_text(encoding="utf-8"))
    sources = [s.strip() for s in args.sources.split(";") if s.strip()]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_summary(ws, body, args.phase, sources, args.summary_note)
    modules: dict[str, list[list[str]]] = {}
    for r in body:
        modules.setdefault(r[2], []).append(r)
    for mod, cases in modules.items():
        write_sheet(wb.create_sheet(mod), cases, notes)
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out} ({len(body)} cases, {len(modules)} module sheets)")


if __name__ == "__main__":
    main()
