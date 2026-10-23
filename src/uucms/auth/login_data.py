"""Load login registry data from Excel when present, otherwise JSON."""

import json
from pathlib import Path
from typing import Any
from zipfile import BadZipFile


def load_login_data(repository_root: Path) -> dict[str, Any]:
    """Load the Excel login registry if present, falling back to the JSON file."""
    login_dir = repository_root / "test_data" / "login"
    excel_path = login_dir / "LoginData.xlsx"
    if excel_path.exists():
        return _load_excel_login_data(excel_path)

    json_path = login_dir / "LoginData.json"
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(
            f"Missing login data: expected {excel_path} or {json_path}"
        ) from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid JSON in login data file: {json_path}") from exc
    if not isinstance(data, dict):
        raise ValueError(f"Login data must contain a JSON object: {json_path}")
    return data


def _load_excel_login_data(excel_path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise ValueError(
            "LoginData.xlsx is present but openpyxl is not installed. "
            "Install project dependencies to read the Excel login registry."
        ) from exc

    try:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
    except (InvalidFileException, OSError, BadZipFile) as exc:
        raise ValueError(f"Unable to read Excel login data file: {excel_path}") from exc

    try:
        worksheet = workbook["Roles"] if "Roles" in workbook.sheetnames else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError(f"Excel login data sheet is empty: {excel_path}")

        normalized_headers = [
            str(header).strip().casefold() if header is not None else ""
            for header in headers
        ]
        header_indexes = {
            header: index
            for index, header in enumerate(normalized_headers)
            if header
        }
        if "role_key" not in header_indexes or "username" not in header_indexes:
            raise ValueError(
                "Excel login data must have 'role_key' and 'username' columns"
            )

        roles: dict[str, dict[str, str]] = {}
        default_account: dict[str, str] | None = None
        for row_number, row in enumerate(rows, start=2):
            if not row or not any(value is not None and str(value).strip() for value in row):
                continue
            role_key = _cell_text(row, header_indexes["role_key"])
            username = _cell_text(row, header_indexes["username"])
            if not role_key or not username:
                raise ValueError(
                    f"Excel login data row {row_number} requires role_key and username"
                )
            if role_key.casefold() == "default" and default_account:
                raise ValueError("Duplicate default login row in Excel data")
            if role_key in roles:
                raise ValueError(f"Duplicate login role key {role_key!r} in Excel data")

            account = {
                "username": username,
                "password_env": _cell_text(
                    row,
                    header_indexes.get("password_env", -1),
                )
                or "UUCMS_TEST_PASSWORD",
            }
            if role_key.casefold() == "default":
                default_account = account
            else:
                roles[role_key] = account

        if default_account is None:
            raise ValueError(
                "Excel login data must include a row with role_key 'default'"
            )

        return {
            **default_account,
            "roles": roles,
        }
    finally:
        workbook.close()


def _cell_text(row: tuple[object, ...], index: int) -> str:
    if index < 0 or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()
